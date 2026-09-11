import os
import re
import markdown
from fpdf import FPDF
from docx import Document
from docx.shared import Pt, Inches

DOCS_DIR = os.path.dirname(os.path.abspath(__file__))


class MarkdownPDF(FPDF):
    def __init__(self):
        super().__init__()
        self.add_font('dejavu', '', 'C:/Windows/Fonts/segoeui.ttf', uni=True)
        self.add_font('dejavu', 'B', 'C:/Windows/Fonts/segoeuib.ttf', uni=True)
        self.add_font('mono', '', 'C:/Windows/Fonts/consola.ttf', uni=True)
        self.set_auto_page_break(auto=True, margin=20)

    def header(self):
        self.set_font('dejavu', '', 8)
        self.set_text_color(150, 150, 150)
        self.cell(0, 8, self.title or '', align='C')
        self.ln(10)
        self.set_draw_color(200, 200, 200)
        self.line(20, self.get_y(), self.w - 20, self.get_y())
        self.ln(4)

    def footer(self):
        self.set_y(-15)
        self.set_font('dejavu', '', 8)
        self.set_text_color(150, 150, 150)
        self.cell(0, 10, f'Página {self.page_no()}/{{nb}}', align='C')

    def write_md(self, md_text):
        lines = md_text.split('\n')
        in_code = False
        code_buf = []
        in_table = False
        table_rows = []

        i = 0
        while i < len(lines):
            line = lines[i]

            # Code blocks
            if line.strip().startswith('```'):
                if in_code:
                    self._flush_code('\n'.join(code_buf))
                    code_buf = []
                    in_code = False
                else:
                    in_code = True
                i += 1
                continue

            if in_code:
                code_buf.append(line)
                i += 1
                continue

            # Table detection
            if '|' in line and line.strip().startswith('|'):
                stripped = line.strip()
                if re.match(r'^\|[\s\-:|]+\|$', stripped):
                    i += 1
                    continue
                cells = [c.strip() for c in stripped.split('|')[1:-1]]
                table_rows.append(cells)
                if i + 1 < len(lines) and '|' in lines[i + 1] and lines[i + 1].strip().startswith('|'):
                    i += 1
                    continue
                else:
                    self._flush_table(table_rows)
                    table_rows = []
                    i += 1
                    continue

            # HR
            if line.strip() == '---':
                self.ln(4)
                self.set_draw_color(180, 180, 180)
                y = self.get_y()
                self.line(20, y, self.w - 20, y)
                self.ln(6)
                i += 1
                continue

            # Headings
            if line.startswith('# ') and not line.startswith('## '):
                self._flush_table(table_rows)
                table_rows = []
                self.ln(4)
                self.set_font('dejavu', 'B', 20)
                self.set_text_color(26, 26, 46)
                self.multi_cell(0, 10, self._strip(line[2:].strip()))
                self.set_draw_color(26, 26, 46)
                y = self.get_y()
                self.line(20, y, self.w - 20, y)
                self.ln(6)
                i += 1
                continue
            if line.startswith('## '):
                self._flush_table(table_rows)
                table_rows = []
                self.ln(3)
                self.set_font('dejavu', 'B', 15)
                self.set_text_color(22, 33, 62)
                self.multi_cell(0, 9, self._strip(line[3:].strip()))
                self.ln(2)
                i += 1
                continue
            if line.startswith('### '):
                self._flush_table(table_rows)
                table_rows = []
                self.ln(2)
                self.set_font('dejavu', 'B', 12)
                self.set_text_color(15, 52, 96)
                self.multi_cell(0, 8, self._strip(line[4:].strip()))
                self.ln(1)
                i += 1
                continue
            if line.startswith('#### '):
                self._flush_table(table_rows)
                table_rows = []
                self.set_font('dejavu', 'B', 11)
                self.set_text_color(15, 52, 96)
                self.multi_cell(0, 7, self._strip(line[5:].strip()))
                self.ln(1)
                i += 1
                continue

            # Empty
            if line.strip() == '':
                self._flush_table(table_rows)
                table_rows = []
                self.ln(2)
                i += 1
                continue

            # List items
            if re.match(r'^(\s*)[-*] ', line):
                self._flush_table(table_rows)
                table_rows = []
                indent = len(line) - len(line.lstrip())
                level = indent // 2
                text = re.sub(r'^(\s*)[-*] ', '', line)
                x_off = 22 + level * 6
                self.set_font('dejavu', '', 10)
                self.set_text_color(34, 34, 34)
                self.set_x(x_off)
                self.cell(4, 6, '•')
                self.set_x(x_off + 5)
                self._write_rich(text, 10)
                self.ln(7)
                i += 1
                continue

            # Numbered list
            if re.match(r'^(\s*)\d+\. ', line):
                self._flush_table(table_rows)
                table_rows = []
                m = re.match(r'^(\s*)(\d+)\. ', line)
                num = m.group(2)
                text = line[m.end():]
                self.set_font('dejavu', '', 10)
                self.set_text_color(34, 34, 34)
                self.set_x(22)
                self.cell(8, 6, f'{num}.')
                self.set_x(30)
                self._write_rich(text, 10)
                self.ln(7)
                i += 1
                continue

            # Regular text
            self._flush_table(table_rows)
            table_rows = []
            self.set_font('dejavu', '', 10)
            self.set_text_color(34, 34, 34)
            self._write_rich(line, 10)
            self.ln(6)
            i += 1

        self._flush_table(table_rows)

    def _strip(self, text):
        text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
        text = re.sub(r'\*(.+?)\*', r'\1', text)
        text = re.sub(r'`(.+?)`', r'\1', text)
        return text

    def _write_rich(self, text, size):
        parts = re.split(r'(\*\*.*?\*\*|`.*?`)', text)
        for part in parts:
            if part.startswith('**') and part.endswith('**'):
                self.set_font('dejavu', 'B', size)
                self.write(6, part[2:-2])
                self.set_font('dejavu', '', size)
            elif part.startswith('`') and part.endswith('`'):
                self.set_font('mono', '', size - 1)
                self.set_text_color(80, 80, 80)
                self.write(6, part[1:-1])
                self.set_font('dejavu', '', size)
                self.set_text_color(34, 34, 34)
            else:
                self.write(6, part)

    def _flush_code(self, code):
        if not code.strip():
            return
        self.set_fill_color(245, 245, 245)
        self.set_font('mono', '', 8.5)
        self.set_text_color(50, 50, 50)
        x = self.get_x()
        self.set_x(22)
        self.multi_cell(self.w - 44, 5, code, fill=True)
        self.ln(4)
        self.set_text_color(34, 34, 34)

    def _flush_table(self, rows):
        if not rows:
            return
        n_cols = max(len(r) for r in rows)
        col_w = (self.w - 40) / n_cols

        for ri, row in enumerate(rows):
            if self.get_y() > self.h - 30:
                self.add_page()
            max_h = 6
            for ci, cell in enumerate(row):
                if ci < n_cols:
                    self.set_xy(20 + ci * col_w, self.get_y())
            if ri == 0:
                self.set_fill_color(232, 232, 232)
                self.set_font('dejavu', 'B', 8.5)
            else:
                self.set_fill_color(255, 255, 255)
                self.set_font('dejavu', '', 8.5)
            self.set_text_color(34, 34, 34)
            self.set_draw_color(200, 200, 200)

            y_start = self.get_y()
            x_start = 20
            max_y = y_start
            cell_texts = []
            for ci in range(n_cols):
                ct = row[ci] if ci < len(row) else ''
                cell_texts.append(self._strip(ct))

            for ci, ct in enumerate(cell_texts):
                self.set_xy(x_start + ci * col_w, y_start)
                self.cell(col_w, 6, ct, border=1, fill=(ri == 0))
            self.ln(6)
        self.ln(4)


def convert_pdf(md_path, pdf_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        md_text = f.read()

    title = ''
    first_line = md_text.split('\n')[0].strip()
    if first_line.startswith('# '):
        title = first_line[2:].strip()

    pdf = MarkdownPDF()
    pdf.set_title(title)
    pdf.alias_nb_pages()
    pdf.add_page()
    pdf.write_md(md_text)
    pdf.output(pdf_path)
    print(f"  PDF: {pdf_path}")


def strip_md_inline(text):
    text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
    text = re.sub(r'\*(.+?)\*', r'\1', text)
    text = re.sub(r'`(.+?)`', r'\1', text)
    return text


def convert_docx(md_path, docx_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.read().split('\n')

    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(4)

    in_code = False
    code_lines = []
    table_rows = []

    i = 0
    while i < len(lines):
        line = lines[i]

        if line.strip().startswith('```'):
            if in_code:
                code_text = '\n'.join(code_lines)
                if code_text.strip():
                    p = doc.add_paragraph()
                    run = p.add_run(code_text)
                    run.font.name = 'Consolas'
                    run.font.size = Pt(9)
                    p.paragraph_format.left_indent = Inches(0.3)
                code_lines = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue

        if in_code:
            code_lines.append(line)
            i += 1
            continue

        if '|' in line and line.strip().startswith('|'):
            stripped = line.strip()
            if re.match(r'^\|[\s\-:|]+\|$', stripped):
                i += 1
                continue
            cells = [c.strip() for c in stripped.split('|')[1:-1]]
            table_rows.append(cells)
            if i + 1 < len(lines) and '|' in lines[i + 1] and lines[i + 1].strip().startswith('|'):
                i += 1
                continue
            else:
                if table_rows:
                    table = doc.add_table(rows=len(table_rows), cols=len(table_rows[0]))
                    table.style = 'Table Grid'
                    for ri, row_data in enumerate(table_rows):
                        for ci, cell_text in enumerate(row_data):
                            if ci < len(table.columns):
                                cell = table.cell(ri, ci)
                                cell.text = strip_md_inline(cell_text)
                                if ri == 0:
                                    for paragraph in cell.paragraphs:
                                        for run in paragraph.runs:
                                            run.bold = True
                    table_rows = []
                i += 1
                continue

        if line.strip() == '---':
            doc.add_paragraph('─' * 60)
            i += 1
            continue

        if line.startswith('# ') and not line.startswith('## '):
            doc.add_heading(strip_md_inline(line[2:].strip()), level=1)
            i += 1
            continue
        if line.startswith('## '):
            doc.add_heading(strip_md_inline(line[3:].strip()), level=2)
            i += 1
            continue
        if line.startswith('### '):
            doc.add_heading(strip_md_inline(line[4:].strip()), level=3)
            i += 1
            continue
        if line.startswith('#### '):
            doc.add_heading(strip_md_inline(line[5:].strip()), level=4)
            i += 1
            continue

        if line.strip() == '':
            i += 1
            continue

        if re.match(r'^(\s*)[-*] ', line):
            text = re.sub(r'^(\s*)[-*] ', '', line)
            p = doc.add_paragraph(style='List Bullet')
            p.clear()
            parts = re.split(r'(\*\*.*?\*\*)', text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                else:
                    p.add_run(strip_md_inline(part))
            i += 1
            continue

        if re.match(r'^(\s*)\d+\. ', line):
            text = re.sub(r'^(\s*)\d+\. ', '', line)
            p = doc.add_paragraph(style='List Number')
            p.clear()
            parts = re.split(r'(\*\*.*?\*\*)', text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                else:
                    p.add_run(strip_md_inline(part))
            i += 1
            continue

        p = doc.add_paragraph()
        parts = re.split(r'(\*\*.*?\*\*|`.*?`)', line)
        for part in parts:
            if part.startswith('**') and part.endswith('**'):
                run = p.add_run(part[2:-2])
                run.bold = True
            elif part.startswith('`') and part.endswith('`'):
                run = p.add_run(part[1:-1])
                run.font.name = 'Consolas'
                run.font.size = Pt(9.5)
            else:
                p.add_run(strip_md_inline(part))
        i += 1

    doc.save(docx_path)
    print(f"  DOCX: {docx_path}")


files = ['MANUAL_USUARIO.md', 'DOCUMENTACION_TECNICA.md']

for fname in files:
    md_path = os.path.join(DOCS_DIR, fname)
    base = os.path.splitext(fname)[0]
    pdf_path = os.path.join(DOCS_DIR, f'{base}.pdf')
    docx_path = os.path.join(DOCS_DIR, f'{base}.docx')

    print(f"Convirtiendo {fname}...")
    convert_pdf(md_path, pdf_path)
    convert_docx(md_path, docx_path)

print("\nListo! Archivos generados en docs/")
